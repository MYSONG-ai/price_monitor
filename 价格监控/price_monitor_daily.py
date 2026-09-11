import csv
import html as html_lib
import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
INPUT = next(
    (
        path
        for path in (
            ROOT / "input.xlsx",
            ROOT / "price_monitor_input_43_updated.xlsx",
            ROOT / "price_monitor_input_43.xlsx",
        )
        if path.exists()
    ),
    ROOT / "input.xlsx",
)
DASHBOARD = ROOT / "price_monitor_dashboard.html"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/131 Safari/537.36",
    "Accept-Language": "it-IT,it;q=0.9,en;q=0.8",
}
FEISHU_BASE_URL = os.getenv("FEISHU_BASE_URL", "https://open.feishu.cn").rstrip("/")
FEISHU_READ_RANGE = os.getenv("FEISHU_READ_RANGE", "A1:AZ200")
DATE_HEADER_START_COL = "H"
FINANCE_TERMS = (
    "finanziamento",
    "rate da",
    "tasso zero",
    "tan",
    "taeg",
    "importo totale",
    "credito",
    "monthlyrate",
)


def parse_price(raw):
    raw = raw.replace(".", "").replace(",", ".")
    try:
        value = float(raw)
    except ValueError:
        return None
    return round(value, 2) if 20 < value < 10000 else None


def parse_json_price(raw):
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        value = float(raw)
        return round(value, 2) if 20 < value < 10000 else None
    text = str(raw).strip()
    match = re.search(r"\d{1,3}(?:\.\d{3})*,\d{2}", text)
    if match:
        return parse_price(match.group(0))
    match = re.search(r"\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        value = float(match.group(0))
    except ValueError:
        return None
    return round(value, 2) if 20 < value < 10000 else None


def iter_json_values(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from iter_json_values(child)
    elif isinstance(value, list):
        for child in value:
            yield from iter_json_values(child)


def extract_jsonld_price(html):
    for match in re.finditer(
        r"<script[^>]+type=[\"']application/ld\+json[\"'][^>]*>(.*?)</script>",
        html,
        flags=re.I | re.S,
    ):
        raw = html_lib.unescape(match.group(1)).strip()
        if not raw:
            continue
        try:
            payload = json.loads(raw)
        except json.JSONDecodeError:
            continue
        for item in iter_json_values(payload):
            offers = item.get("offers") if isinstance(item, dict) else None
            if isinstance(offers, list):
                offer_items = offers
            elif isinstance(offers, dict):
                offer_items = [offers]
            else:
                offer_items = []
            for offer in offer_items:
                price = parse_json_price(offer.get("price") or offer.get("lowPrice"))
                if price is not None:
                    return price, "json-ld"
    return None, None


def mediaworld_product_id(url):
    if not url:
        return None
    match = re.search(r"-(\d+)\.html(?:[?#].*)?$", url)
    return match.group(1) if match else None


def extract_mediaworld_cofr_price(html, url):
    product_id = mediaworld_product_id(url)
    if not product_id:
        return None, None
    markers = (f'"id":"Media:it:{product_id}"', f'\\"id\\":\\"Media:it:{product_id}\\"')
    for marker in markers:
        start = 0
        while True:
            index = html.find(marker, start)
            if index < 0:
                break
            segment = html[index:index + 8000]
            match = re.search(
                r'"price"\s*:\s*\{[\s\S]{0,600}?"amount"\s*:\s*([0-9]+(?:\.[0-9]+)?)',
                segment,
            )
            if match:
                price = parse_json_price(match.group(1))
                if price is not None:
                    return price, "mediaworld-cofr-price"
            start = index + len(marker)
    return None, None


def is_unieuro_search_url(url):
    return bool(url and "/online/products" in url and "?" in url and "q=" in url.lower())


def extract_text_price(html, retailer):
    text = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>|<[^>]+>", " ", html, flags=re.I)
    text = re.sub(r"\s+", " ", text.replace("&euro;", "€"))
    patterns = (
        [r"(\d{1,3}(?:\.\d{3})*,\d{2})\s*€\s*(?:CONSEGNA|IVA inclusa)", r"(\d{1,3}(?:\.\d{3})*,\d{2})\s*€"]
        if retailer == "MediaWorld"
        else [r"€\s*(\d{1,3}(?:\.\d{3})*,\d{2})\s*Iva inclusa", r"€\s*(\d{1,3}(?:\.\d{3})*,\d{2})"]
    )
    for pattern in patterns:
        for match in re.finditer(pattern, text, re.I):
            if retailer == "MediaWorld":
                context = text[max(0, match.start() - 160):match.end() + 160].lower()
                if any(term in context for term in FINANCE_TERMS):
                    continue
            price = parse_price(match.group(1))
            if price is not None:
                return price, "text"
    return None, None


def extract_price_detail(html, retailer, url=None):
    if retailer == "Unieuro" and is_unieuro_search_url(url):
        return None, "ignored-unieuro-search-page"

    if retailer == "MediaWorld":
        price, source = extract_mediaworld_cofr_price(html, url)
        if price is not None:
            return price, source

    price, source = extract_jsonld_price(html)
    if price is not None:
        return price, source

    return extract_text_price(html, retailer)


def extract_price(html, retailer, url=None):
    price, _source = extract_price_detail(html, retailer, url)
    return price


def split_item(item):
    normalized = str(item).replace("_", " ").strip()
    match = re.match(r"^([A-Za-z]+)\s+(\d{2})(.*)$", normalized)
    if not match:
        return "", "", normalized
    return match.group(2), match.group(1).upper(), f"{match.group(2)}{match.group(3)}"


def derive_size(*values):
    text = " ".join(str(value or "") for value in values)
    match = re.search(r"(?<!\d)(32|43|50|55|65|75|85)(?=\D|$)", text)
    return match.group(1) if match else ""


def normalize_header(value):
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def find_column(headers, names=(), contains=(), required=True):
    normalized_names = {normalize_header(name) for name in names}
    normalized_contains = [normalize_header(term) for term in contains]
    for index, header in enumerate(headers):
        normalized = normalize_header(header)
        if normalized in normalized_names:
            return index
        if any(term and term in normalized for term in normalized_contains):
            return index
    if required:
        expected = ", ".join([*names, *contains])
        raise RuntimeError(f"{INPUT.name} is missing required header: {expected}")
    return None


def row_value(row, index):
    if index is None or index >= len(row):
        return ""
    value = row[index]
    return str(value).strip() if value not in (None, "") else ""


def normalize_key_part(value):
    return re.sub(r"\s+", "", str(value or "")).upper()


def price_to_cell(value):
    if isinstance(value, (int, float)):
        return f"{value:g}"
    return ""


def col_to_num(col):
    total = 0
    for char in col.upper():
        total = total * 26 + ord(char) - 64
    return total


def num_to_col(num):
    chars = []
    while num:
        num, rem = divmod(num - 1, 26)
        chars.append(chr(65 + rem))
    return "".join(reversed(chars))


def next_col(col):
    return num_to_col(col_to_num(col) + 1)


def ordinal_day(day):
    if 10 <= day % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return f"{day}{suffix}"


def feishu_date_label(date_text):
    date_value = datetime.strptime(date_text, "%Y-%m-%d").date()
    return f"{ordinal_day(date_value.day)} {date_value.strftime('%b')}"


def parse_date_header(value, default_year):
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    match = re.fullmatch(r"(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3,})", text, flags=re.I)
    if match:
        try:
            return datetime.strptime(f"{default_year} {match.group(1)} {match.group(2)[:3]}", "%Y %d %b").date()
        except ValueError:
            return None
    return None


def load_products():
    workbook = load_workbook(INPUT, data_only=True)
    sheet = workbook["Input"] if "Input" in workbook.sheetnames else workbook.active
    header_row = next(sheet.iter_rows(values_only=True))
    headers = [str(value or "").strip() for value in header_row]
    item_col = find_column(headers, names=("item name",), required=False)
    brand_col = find_column(headers, names=("brand", "品牌"), required=False)
    model_col = find_column(headers, names=("model", "型号"), required=False)
    technology_col = find_column(headers, names=("technology", "tech", "技术", "清晰度", "列1"), required=False)
    ue_col = find_column(headers, contains=("unieuro",))
    mw_col = find_column(headers, contains=("mediaworld",))

    if item_col is None and (brand_col is None or model_col is None):
        raise RuntimeError(f"{INPUT.name} must include either Item Name or Brand/Model columns")

    products = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        unieuro_url = row_value(row, ue_col)
        mediaworld_url = row_value(row, mw_col)
        if item_col is not None:
            item = row_value(row, item_col)
            if not item:
                continue
            size, brand, model = split_item(item)
            technology = row_value(row, technology_col)
        else:
            brand = row_value(row, brand_col).upper()
            model = row_value(row, model_col)
            technology = row_value(row, technology_col)
            if not any((brand, model, technology, unieuro_url, mediaworld_url)):
                continue
            if not brand or not model:
                print(f"skipped input row {row_number}; missing brand or model")
                continue
            size = derive_size(model, unieuro_url, mediaworld_url)
            item = f"{brand} {model}"

        if not unieuro_url and not mediaworld_url:
            continue

        products.append({
            "item": item,
            "size": size,
            "brand": brand,
            "model": model,
            "technology": technology,
            "Unieuro": unieuro_url or None,
            "MediaWorld": mediaworld_url or None,
        })
    return products


def fetch_price(url, retailer, attempts=3):
    if not url:
        return "/"
    for attempt in range(1, attempts + 1):
        try:
            request = Request(url, headers=HEADERS)
            with urlopen(request, timeout=30) as response:
                price, source = extract_price_detail(response.read().decode("utf-8", errors="ignore"), retailer, url)
            if price is not None:
                if os.getenv("PRICE_MONITOR_DEBUG", "").lower() in {"1", "true", "yes", "on"}:
                    print(f"{retailer} {url} -> {price} ({source})")
                return price
            if os.getenv("PRICE_MONITOR_DEBUG", "").lower() in {"1", "true", "yes", "on"}:
                print(f"{retailer} {url} -> no price ({source or 'not-found'}), attempt {attempt}/{attempts}")
        except (HTTPError, URLError, TimeoutError) as error:
            if os.getenv("PRICE_MONITOR_DEBUG", "").lower() in {"1", "true", "yes", "on"}:
                print(f"{retailer} {url} -> fetch error {type(error).__name__}, attempt {attempt}/{attempts}")
        if attempt < attempts:
            time.sleep(2 * attempt)
    return None


def env_truthy(name):
    return os.getenv(name, "").lower() in {"1", "true", "yes", "on"}


def feishu_writes_enabled():
    return env_truthy("FEISHU_WRITE")


def feishu_required_env():
    names = ("FEISHU_APP_ID", "FEISHU_APP_SECRET", "FEISHU_SPREADSHEET_TOKEN", "FEISHU_SHEET_ID")
    values = {name: os.getenv(name) for name in names}
    missing = [name for name, value in values.items() if not value]
    return values, missing


def feishu_access_token(app_id, app_secret):
    response = requests.post(
        f"{FEISHU_BASE_URL}/open-apis/auth/v3/tenant_access_token/internal",
        json={"app_id": app_id, "app_secret": app_secret},
        timeout=60,
    )
    response.raise_for_status()
    payload = response.json()
    if payload.get("code", 0) != 0:
        raise RuntimeError(f"Feishu auth failed: {payload.get('msg') or payload}")
    return payload["tenant_access_token"]


def feishu_invoke_tool(access_token, spreadsheet_token, tool_name, tool_input):
    response = requests.post(
        f"{FEISHU_BASE_URL}/open-apis/sheet_ai/v2/spreadsheets/{spreadsheet_token}/tools/invoke_write"
        if tool_name in {"set_range_from_csv", "modify_sheet_structure"}
        else f"{FEISHU_BASE_URL}/open-apis/sheet_ai/v2/spreadsheets/{spreadsheet_token}/tools/invoke_read",
        headers={"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"},
        json={"tool_name": tool_name, "input": json.dumps(tool_input, ensure_ascii=False)},
        timeout=90,
    )
    response.raise_for_status()
    payload = response.json()
    if payload.get("code", 0) != 0:
        raise RuntimeError(f"Feishu tool {tool_name} failed: {payload.get('msg') or payload}")
    data = payload.get("data") or {}
    output = data.get("output") or data.get("result") or data
    if isinstance(output, str):
        try:
            return json.loads(output)
        except json.JSONDecodeError:
            return {"output": output}
    return output


def feishu_read_rows(access_token, spreadsheet_token, sheet_id):
    result = feishu_invoke_tool(access_token, spreadsheet_token, "get_range_as_csv", {
        "excel_id": spreadsheet_token,
        "sheet_id": sheet_id,
        "range": FEISHU_READ_RANGE,
        "max_chars": 200000,
        "max_rows": 1000000000,
    })
    if "rows" in result:
        return result["rows"]
    annotated_csv = result.get("annotated_csv") or result.get("output") or ""
    rows = []
    for index, line in enumerate(annotated_csv.splitlines(), start=1):
        row_number = index
        match = re.match(r"^\[row=(\d+)\]\s*(.*)$", line)
        if match:
            row_number = int(match.group(1))
            line = match.group(2)
        values = next(csv.reader([line]))
        rows.append({"row_number": row_number, "values": {num_to_col(i + 1): value for i, value in enumerate(values)}})
    return rows


def feishu_write_csv(access_token, spreadsheet_token, sheet_id, start_cell, csv_text):
    return feishu_invoke_tool(access_token, spreadsheet_token, "set_range_from_csv", {
        "excel_id": spreadsheet_token,
        "sheet_id": sheet_id,
        "start_cell": start_cell,
        "csv": csv_text,
    })


def feishu_insert_column(access_token, spreadsheet_token, sheet_id, column):
    return feishu_invoke_tool(access_token, spreadsheet_token, "modify_sheet_structure", {
        "excel_id": spreadsheet_token,
        "sheet_id": sheet_id,
        "operation": "insert",
        "position": column,
        "count": 1,
        "side": "before",
    })


def header_lookup(header_row):
    aliases = {
        "size": {"size", "尺寸"},
        "brand": {"brand", "品牌"},
        "model": {"model", "型号"},
        "channel": {"channel", "渠道"},
        "link": {"link", "链接"},
    }
    by_name = {}
    for col, value in (header_row.get("values") or {}).items():
        normalized = str(value or "").strip().lower()
        for key, names in aliases.items():
            if normalized in {name.lower() for name in names}:
                by_name[key] = col
    required = {"size", "brand", "model", "channel"}
    missing = sorted(required - by_name.keys())
    if missing:
        raise RuntimeError(f"Feishu sheet is missing required headers: {', '.join(missing)}")
    return by_name


def find_or_create_date_column(access_token, spreadsheet_token, sheet_id, header_row, today):
    values = header_row.get("values") or {}
    today_date = datetime.strptime(today, "%Y-%m-%d").date()
    today_labels = {today, feishu_date_label(today)}
    for col, value in values.items():
        if str(value or "").strip() in today_labels:
            return col

    date_cols = []
    start_num = col_to_num(DATE_HEADER_START_COL)
    for col, value in values.items():
        if col_to_num(col) < start_num:
            continue
        parsed = parse_date_header(value, today_date.year)
        if parsed:
            date_cols.append((parsed, col))

    previous = [item for item in date_cols if item[0] < today_date]
    if previous:
        insert_at = next_col(max(previous, key=lambda item: item[0])[1])
    elif date_cols:
        insert_at = min(date_cols, key=lambda item: item[0])[1]
    else:
        insert_at = DATE_HEADER_START_COL

    feishu_insert_column(access_token, spreadsheet_token, sheet_id, insert_at)
    feishu_write_csv(access_token, spreadsheet_token, sheet_id, f"{insert_at}1", feishu_date_label(today))
    return insert_at


def update_feishu_sheet(current, today):
    if not feishu_writes_enabled():
        print("skipped Feishu update; set FEISHU_WRITE=true to enable")
        return

    env, missing = feishu_required_env()
    if missing:
        print(f"skipped Feishu update; missing secrets: {', '.join(missing)}")
        return

    access_token = feishu_access_token(env["FEISHU_APP_ID"], env["FEISHU_APP_SECRET"])
    spreadsheet_token = env["FEISHU_SPREADSHEET_TOKEN"]
    sheet_id = env["FEISHU_SHEET_ID"]
    rows = feishu_read_rows(access_token, spreadsheet_token, sheet_id)
    if not rows:
        raise RuntimeError("Feishu sheet is empty")

    columns = header_lookup(rows[0])
    date_col = find_or_create_date_column(access_token, spreadsheet_token, sheet_id, rows[0], today)
    prices = {
        (
            normalize_key_part(row["size"]),
            normalize_key_part(row["brand"]),
            normalize_key_part(row["model"]),
            normalize_key_part(row["channel"]),
        ): row["price"]
        for row in current
        if row.get("price") not in (None, "/")
    }

    last_size = last_brand = last_model = ""
    writes = 0
    for row in rows[1:]:
        values = row.get("values") or {}
        size = values.get(columns["size"], "")
        brand = values.get(columns["brand"], "")
        model = values.get(columns["model"], "")
        channel = values.get(columns["channel"], "")
        if size:
            last_size = size
        if brand:
            last_brand = brand
        if model:
            last_model = model
        if not channel:
            continue
        key = (
            normalize_key_part(last_size),
            normalize_key_part(last_brand),
            normalize_key_part(last_model),
            normalize_key_part(channel),
        )
        if key not in prices:
            continue
        feishu_write_csv(access_token, spreadsheet_token, sheet_id, f"{date_col}{row['row_number']}", price_to_cell(prices[key]))
        writes += 1
    print(f"updated Feishu sheet column {date_col} with {writes} prices for {today}")


def load_history(html, fallback_date):
    match = re.search(r"const DATA = (\[.*?\]);\s*\nconst \$", html, flags=re.S)
    if not match:
        match = re.search(r"const DATA = (\[.*?\])\.map", html, flags=re.S)
    history = json.loads(match.group(1)) if match else []
    existing_date_match = re.search(r"最新数据：([0-9]{4}-[0-9]{2}-[0-9]{2})", html)
    existing_date = existing_date_match.group(1) if existing_date_match else fallback_date
    for row in history:
        row.setdefault("date", existing_date)
    return history


def main():
    today = datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat()
    html = DASHBOARD.read_text(encoding="utf-8")
    history = load_history(html, today)
    products = load_products()
    current = []
    valid_keys = set()
    tasks = []
    for product in products:
        size = product.get("size")
        brand = product.get("brand")
        model = product.get("model")
        technology = product.get("technology", "")
        if not (size and brand and model):
            size, brand, model = split_item(product["item"])
        for channel in ("Unieuro", "MediaWorld"):
            if not product[channel]:
                continue
            valid_keys.add((size, brand, model, channel))
            tasks.append((size, brand, model, technology, channel, product[channel]))

    def scrape(task):
        size, brand, model, technology, channel, url = task
        return {
            "size": size,
            "brand": brand,
            "model": model,
            "technology": technology,
            "channel": channel,
            "price": fetch_price(url, channel),
            "date": today,
        }

    with ThreadPoolExecutor(max_workers=12) as pool:
        current = list(pool.map(scrape, tasks))

    missing = [index for index, row in enumerate(current) if row.get("price") is None]
    if missing:
        print(f"retrying {len(missing)} missing prices in sequential fallback")
        for index in missing:
            size, brand, model, technology, channel, url = tasks[index]
            retry_price = fetch_price(url, channel, attempts=3)
            if retry_price is not None:
                current[index] = {
                    "size": size,
                    "brand": brand,
                    "model": model,
                    "technology": technology,
                    "channel": channel,
                    "price": retry_price,
                    "date": today,
                }

    keys = {(r["size"], r["brand"], r["model"], r["channel"]) for r in current}
    history = [r for r in history if (r.get("size"), r.get("brand"), r.get("model"), r.get("channel")) in valid_keys and ((r.get("size"), r.get("brand"), r.get("model"), r.get("channel")) not in keys or r.get("date") != today)]
    payload = json.dumps(history + current, ensure_ascii=False, separators=(",", ":"))
    html = re.sub(r"const DATA = .*?;\s*\nconst \$", f"const DATA = {payload};\nconst $", html, count=1, flags=re.S)
    html = re.sub(r"最新数据：[0-9]{4}-[0-9]{2}-[0-9]{2}", f"最新数据：{today}", html)
    if env_truthy("PRICE_MONITOR_DRY_RUN"):
        missing_count = sum(1 for row in current if row.get("price") is None)
        print(f"dry run: fetched {len(current)} records for {today}; {missing_count} missing; skipped dashboard and Feishu writes")
        return
    DASHBOARD.write_text(html, encoding="utf-8")
    print(f"updated {DASHBOARD} with {len(current)} records for {today}")
    update_feishu_sheet(current, today)


if __name__ == "__main__":
    main()
